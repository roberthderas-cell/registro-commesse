"""Connettore per file Excel (.xlsx) e CSV: estrazioni dal gestionale, budget, ordini, rapportini, incassi."""
import csv
import io

from openpyxl import load_workbook

from .base import TIPI, Connettore, Lettura, Record, nrm

# Nomi di colonna alternativi che si trovano spesso nelle estrazioni dei gestionali
ALIAS = {
    "cod": ["codice", "codicecommessa", "commessa", "nrcommessa", "job"],
    "commessa": ["codicecommessa", "cod", "nrcommessa", "job", "centrodicosto", "cdc"],
    "numero": ["num", "nfattura", "numerofattura", "nrfattura", "numdoc", "numerodocumento"],
    "fornitore": ["ragionesocialefornitore", "denominazionefornitore", "fornitoreragionesociale"],
    "cliente": ["ragionesociale", "ragionesocialecliente", "denominazione"],
    "importo": ["imponibile", "totaleimponibile", "valore", "ammontare", "importoeur"],
    "data": ["datadocumento", "datafattura", "dataordine", "datacontabile", "dataoperazione", "datavaluta"],
    "categoria": ["voce", "vocedicosto", "categoriacosto", "famiglia"],
    "ordine": ["nordine", "numeroordine", "rifordine", "oda"],
    "descrizione": ["descr", "oggetto", "articolo"],
    "avanzamento": ["avanz", "percentuale", "avanzamentopercentuale", "sal"],
    "ore": ["oreuomo", "orelavorate"],
    "chi": ["operatore", "capocantiere", "autore"],
    "causale": ["descrizioneoperazione", "descrizionemovimento", "dettaglio"],
    "ricavo": ["importocontratto", "valorecontratto"],
    "contratto": ["datacontratto", "firma"],
    "scadenza": ["datascadenza"],
    "piva_fornitore": ["partitaivafornitore", "pivafornitore"],
    "piva_cliente": ["partitaivacliente", "pivacliente"],
}


def leggi_righe(contenuto: bytes, nome_file: str) -> list[list]:
    if nome_file.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    testo = contenuto.decode("utf-8-sig", errors="replace")
    prima = testo.split("\n", 1)[0]
    sep = "\t" if "\t" in prima else (";" if prima.count(";") >= prima.count(",") else ",")
    return [r for r in csv.reader(io.StringIO(testo), delimiter=sep) if any(c.strip() for c in r)]


def mappa_colonne(intestazioni: list, campi: list[str]) -> dict[str, int]:
    hd = [nrm(h) for h in intestazioni]
    ix = {}
    for campo in campi:
        for nome in [nrm(campo), *ALIAS.get(campo, [])]:
            if nome in hd:
                ix[campo] = hd.index(nome)
                break
    return ix


class TabelleConnettore(Connettore):
    nome = "tabelle"
    etichetta = "Excel o CSV"
    descrizione = "File .xlsx o .csv con una riga di intestazioni: commesse, ordini, fatture, rapportini, incassi."
    estensioni = [".xlsx", ".xlsm", ".csv", ".txt"]
    tipi = ["commessa", "ordine", "fattura_passiva", "fattura_attiva", "rapportino", "movimento"]
    serve_tipo = True

    def leggi(self, contenuto: bytes, nome_file: str, tipo: str | None = None) -> Lettura:
        out = Lettura()
        if tipo not in self.tipi:
            out.errore("file", f"Indica che cosa contiene il file ({', '.join(self.tipi)})")
            return out
        try:
            righe = leggi_righe(contenuto, nome_file)
        except Exception as e:  # file corrotto o formato sconosciuto
            out.errore("file", f"Impossibile leggere il file: {e}")
            return out
        if len(righe) < 2:
            out.errore("file", "Servono la riga delle intestazioni e almeno una riga di dati")
            return out
        spec = TIPI[tipo]
        ix = mappa_colonne(righe[0], spec["req"] + spec["opt"])
        mancano = [k for k in spec["req"] if k not in ix]
        if mancano:
            out.errore("riga 1", "Colonne mancanti nelle intestazioni: " + ", ".join(mancano))
            return out
        for n, r in enumerate(righe[1:], start=2):
            dati = {k: (r[i] if i < len(r) else None) for k, i in ix.items()}
            dati = {k: (v.strip() if isinstance(v, str) else v) for k, v in dati.items()}
            out.records.append(Record(tipo, dati, f"riga {n}", rif_esterno=_rif(tipo, dati)))
        return out


def _rif(tipo: str, d: dict) -> str | None:
    if tipo == "ordine" and d.get("id"):
        return str(d["id"])
    if tipo in ("fattura_passiva", "fattura_attiva") and d.get("numero"):
        return f"{d.get('fornitore') or d.get('cliente') or ''}|{d['numero']}"
    if tipo == "movimento" and d.get("rif"):
        return str(d["rif"])
    return None
